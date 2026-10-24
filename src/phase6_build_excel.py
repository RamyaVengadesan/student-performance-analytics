"""
Phase 6a: build the Excel dashboard workbook (formula-driven).

Usage:  python phase6_build_excel.py   (then recalc with LibreOffice if you want cached values)
Needs:  data/processed/students_clean.csv, data/processed/student_risk_scores.csv
Output: excel/Student_Performance_Dashboard.xlsx, data/processed/powerbi_students.csv
"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY, BLUE, LIGHT, GREY = "1F2937", "2563EB", "EFF6FF", "F3F4F6"
OUT = "excel/Student_Performance_Dashboard.xlsx"
import os; os.makedirs("excel", exist_ok=True)

df = pd.read_csv("data/processed/students_clean.csv", dtype={"absence_band": str})
df.insert(0, "student_id", range(1, len(df) + 1))
risk = pd.read_csv("data/processed/student_risk_scores.csv")[["student_id", "risk_of_failing", "risk_level"]]
df = df.merge(risk, on="student_id", how="left")
# Labels that Excel cannot mistake for dates (a criterion like "1-5" would be read as a date)
band_map = {"0": "None (0)", "1-5": "Low (1-5)", "6-10": "Medium (6-10)", "11+": "High (11+)"}
df["absence_band"] = df["absence_band"].map(band_map)
df["higher_label"] = df["higher"].map({1: "Yes", 0: "No"})
df.to_csv("data/processed/powerbi_students.csv", index=False)

N = len(df)
last = N + 1
col = {name: get_column_letter(i + 1) for i, name in enumerate(df.columns)}


def rng(name):
    return f"Data!${col[name]}$2:${col[name]}${last}"


wb = Workbook()
thin = Side(style="thin", color="D1D5DB")
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, c1, c2):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = box


# ------------------------------------------------------------------ Data sheet
ws_d = wb.active
ws_d.title = "Data"
ws_d.append(list(df.columns))
for row in df.itertuples(index=False):
    ws_d.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])
style_header(ws_d, 1, 1, len(df.columns))
for c in range(1, len(df.columns) + 1):
    ws_d.column_dimensions[get_column_letter(c)].width = 13
    for r in range(2, last + 1):
        ws_d.cell(row=r, column=c).font = Font(name=FONT, size=10)
ws_d.freeze_panes = "B2"
ws_d.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{last}"
for r in range(2, last + 1):
    ws_d[f"{col['risk_of_failing']}{r}"].number_format = "0.0%"

# ------------------------------------------------------------------ Summary sheet
ws_s = wb.create_sheet("Summary")
ws_s.sheet_view.showGridLines = False
ws_s["A1"] = "Summary tables (all values are formulas on the Data sheet)"
ws_s["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
ws_s["A2"] = "Pass = final grade G3 of 10 or more. Tables marked 'active' exclude the 53 records with G3 = 0 (possible dropouts, no attendance data)."
ws_s["A2"].font = Font(name=FONT, size=9, italic=True, color="6B7280")

DROP = rng("dropout_flag")
G3, PASS, ABS = rng("G3"), rng("pass_flag"), rng("absences")
tables = {}


def table(title, top, headers, labels, key_range, make_formulas, key_is_text=True, note=None):
    ws_s.cell(row=top, column=1, value=title).font = Font(name=FONT, bold=True, size=11, color=NAVY)
    if note:
        ws_s.cell(row=top, column=len(headers) + 1, value=note).font = Font(name=FONT, size=9, italic=True, color="6B7280")
    for i, h in enumerate(headers):
        ws_s.cell(row=top + 1, column=1 + i, value=h)
    style_header(ws_s, top + 1, 1, len(headers))
    for j, lab in enumerate(labels):
        r = top + 2 + j
        ws_s.cell(row=r, column=1, value=lab)
        for k, f in enumerate(make_formulas(r), start=2):
            ws_s.cell(row=r, column=k, value=f)
        for c in range(1, len(headers) + 1):
            cell = ws_s.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10); cell.border = box
    tables[title] = (top + 1, top + 1 + len(labels))
    return top + len(labels) + 4


fmt = {"pct": "0.0%", "avg": "0.00", "int": "#,##0"}

# 1 Subject
t1 = 4
nxt = table("1. Results by subject (all records)", t1,
            ["Subject", "Records", "Avg G3", "Pass rate", "Avg absences"], ["Maths", "Portuguese"],
            rng("subject"),
            lambda r: [f"=COUNTIFS({rng('subject')},$A{r})",
                       f"=AVERAGEIFS({G3},{rng('subject')},$A{r})",
                       f"=COUNTIFS({rng('subject')},$A{r},{PASS},1)/B{r}",
                       f"=AVERAGEIFS({ABS},{rng('subject')},$A{r})"])
# 2 Absence band (active)
t2 = nxt
bands = list(band_map.values())
nxt = table("2. Absence band vs results (active records)", t2,
            ["Absence band", "Records", "Avg G3", "Pass rate"], bands, rng("absence_band"),
            lambda r: [f"=COUNTIFS({rng('absence_band')},$A{r},{DROP},0)",
                       f"=AVERAGEIFS({G3},{rng('absence_band')},$A{r},{DROP},0)",
                       f"=COUNTIFS({rng('absence_band')},$A{r},{DROP},0,{PASS},1)/B{r}"])
# 3 Failures
t3 = nxt
nxt = table("3. Past failures vs results (all records)", t3,
            ["Past failures", "Records", "Avg G3", "Pass rate"], [0, 1, 2, 3], rng("failures"),
            lambda r: [f"=COUNTIFS({rng('failures')},$A{r})",
                       f"=AVERAGEIFS({G3},{rng('failures')},$A{r})",
                       f"=COUNTIFS({rng('failures')},$A{r},{PASS},1)/B{r}"])
# 4 Higher education
t4 = nxt
nxt = table("4. Wants higher education (all records)", t4,
            ["Wants higher edu", "Records", "Avg G3", "Pass rate"], ["Yes", "No"], rng("higher_label"),
            lambda r: [f"=COUNTIFS({rng('higher_label')},$A{r})",
                       f"=AVERAGEIFS({G3},{rng('higher_label')},$A{r})",
                       f"=COUNTIFS({rng('higher_label')},$A{r},{PASS},1)/B{r}"])
# 5 Study time
t5 = nxt
nxt = table("5. Study time level (active records)", t5,
            ["Study time (1-4)", "Records", "Avg G3", "Avg grade change (G3-G1)"], [1, 2, 3, 4], rng("studytime"),
            lambda r: [f"=COUNTIFS({rng('studytime')},$A{r},{DROP},0)",
                       f"=AVERAGEIFS({G3},{rng('studytime')},$A{r},{DROP},0)",
                       f"=AVERAGEIFS({rng('grade_change')},{rng('studytime')},$A{r},{DROP},0)"])
# 6 Risk level
t6 = nxt
nxt = table("6. Early-warning risk level (model uses G1 and background, cross-validated)", t6,
            ["Risk level", "Records", "Actual fail rate"], ["Low", "Medium", "High"], rng("risk_level"),
            lambda r: [f"=COUNTIFS({rng('risk_level')},$A{r})",
                       f"=COUNTIFS({rng('risk_level')},$A{r},{PASS},0)/B{r}"])

for title, (hdr, endr) in tables.items():
    headers = [ws_s.cell(row=hdr, column=c).value for c in range(1, 8) if ws_s.cell(row=hdr, column=c).value]
    for r in range(hdr + 1, endr + 1):
        for c, h in enumerate(headers, start=1):
            cell = ws_s.cell(row=r, column=c)
            if h == "Records":
                cell.number_format = fmt["int"]
            elif h in ("Pass rate", "Actual fail rate"):
                cell.number_format = fmt["pct"]
            elif h.startswith("Avg"):
                cell.number_format = fmt["avg"]
            if c > 1:
                cell.alignment = Alignment(horizontal="center")
ws_s.column_dimensions["A"].width = 26
for c in "BCDEF":
    ws_s.column_dimensions[c].width = 18

# ------------------------------------------------------------------ Dashboard sheet
ws = wb.create_sheet("Dashboard", 0)
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 2
for c in range(2, 14):
    ws.column_dimensions[get_column_letter(c)].width = 13
ws.merge_cells("B1:M1")
ws["B1"] = "Student Performance & Attendance Analytics"
ws["B1"].font = Font(name=FONT, size=20, bold=True, color="FFFFFF")
ws["B1"].fill = PatternFill("solid", fgColor=NAVY)
ws["B1"].alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[1].height = 38
ws.merge_cells("B2:M2")
ws["B2"] = "UCI Student Performance dataset (Cortez & Silva, 2008) - Maths and Portuguese, 1,044 records from two Portuguese secondary schools"
ws["B2"].font = Font(name=FONT, size=9, italic=True, color="6B7280")

kpis = [
    ("Records (student-subject)", f"=COUNTA({rng('student_id')})", "#,##0"),
    ("Pass rate", f"=SUM({PASS})/COUNTA({rng('student_id')})", "0.0%"),
    ("Avg final grade", f"=AVERAGE({G3})", "0.00"),
    ("Avg absences", f"=AVERAGE({ABS})", "0.00"),
    ("High-risk records", f"=COUNTIFS({rng('risk_level')},\"High\")", "#,##0"),
    ("Possible dropouts (G3=0)", f"=SUM({DROP})", "#,##0"),
]
for i, (label, formula, nf) in enumerate(kpis):
    c1 = 2 + i * 2
    a, b = get_column_letter(c1), get_column_letter(c1 + 1)
    ws.merge_cells(f"{a}4:{b}4"); ws.merge_cells(f"{a}5:{b}5")
    ws[f"{a}4"] = label
    ws[f"{a}4"].font = Font(name=FONT, size=9, bold=True, color="6B7280")
    ws[f"{a}5"] = formula
    ws[f"{a}5"].font = Font(name=FONT, size=20, bold=True, color=BLUE)
    ws[f"{a}5"].number_format = nf
    for r in (4, 5):
        for cc in (a, b):
            ws[f"{cc}{r}"].fill = PatternFill("solid", fgColor=LIGHT)
            ws[f"{cc}{r}"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[5].height = 34


def bar(title, tbl, val_col, y_title, x_title, anchor, pct=False, color=BLUE, ymax=None):
    hdr, endr = tables[tbl]
    ch = BarChart(); ch.type = "col"; ch.style = 10
    ch.title = title
    data = Reference(ws_s, min_col=val_col, min_row=hdr, max_row=endr)
    cats = Reference(ws_s, min_col=1, min_row=hdr + 1, max_row=endr)
    ch.add_data(data, titles_from_data=True); ch.set_categories(cats)
    ch.legend = None
    ch.y_axis.title = y_title; ch.x_axis.title = x_title
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.y_axis.majorGridlines = None
    if pct:
        ch.y_axis.number_format = "0%"; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1
    elif ymax:
        ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = ymax
    ch.series[0].graphicalProperties.solidFill = color
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
    ch.dataLabels.showSerName = False; ch.dataLabels.showCatName = False; ch.dataLabels.showLegendKey = False
    if pct:
        ch.dataLabels.numFmt = "0%"
    else:
        ch.dataLabels.numFmt = "0.0"
    ch.width, ch.height = 15.5, 7.5
    ws.add_chart(ch, anchor)


bar("Pass rate by absence band (active)", "2. Absence band vs results (active records)", 4,
    "Pass rate", "Absences", "B8", pct=True)
bar("Pass rate by past failures", "3. Past failures vs results (all records)", 4,
    "Pass rate", "Past failures", "H8", pct=True, color="DC2626")
bar("Average final grade by subject", "1. Results by subject (all records)", 3,
    "Avg G3 (0-20)", "Subject", "B24", ymax=20, color="F59E0B")
bar("Actual fail rate by predicted risk level", "6. Early-warning risk level (model uses G1 and background, cross-validated)", 3,
    "Actual fail rate", "Risk level", "H24", pct=True, color="7C3AED")

ws["B40"] = "Charts read from the Summary sheet, which is formula-driven from the Data sheet. See Notes for definitions."
ws["B40"].font = Font(name=FONT, size=9, italic=True, color="6B7280")

# ------------------------------------------------------------------ Notes sheet
ws_n = wb.create_sheet("Notes")
ws_n.sheet_view.showGridLines = False
ws_n.column_dimensions["A"].width = 28; ws_n.column_dimensions["B"].width = 110
notes = [
    ("Source", "P. Cortez and A. Silva (2008), Using Data Mining to Predict Secondary School Student Performance. UCI Machine Learning Repository, Student Performance dataset, CC BY 4.0."),
    ("Records", "1,044 rows: 395 Maths and 649 Portuguese. The two files are stacked, so many students appear twice. There is no student id; matching on 13 background attributes finds about 366 shared profiles, which is approximate."),
    ("Pass", "pass_flag = 1 when the final grade G3 is 10 or more (grades are on a 0-20 scale)."),
    ("Possible dropouts", "dropout_flag = 1 when G3 = 0. These 53 records all show 0 absences, so attendance was probably not recorded. Tables marked 'active' exclude them."),
    ("Absence bands", "None (0), Low (1-5), Medium (6-10), High (11+). Relabelled so Excel does not read '1-5' as a date."),
    ("Risk score", "Out-of-fold probability of failing from a Phase 5 model using G1 and background factors (not G2 or G3). Folds are grouped so a student's Maths and Portuguese rows are never split between training and testing. Low < 20%, Medium 20-50%, High > 50%."),
    ("Association, not cause", "All findings are associations from observational data."),
    ("To refresh", "Paste new rows into the Data sheet (same columns). Dashboard and Summary recalculate. Ranges cover rows 2 to %d." % last),
]
ws_n["A1"] = "Notes and assumptions"; ws_n["A1"].font = Font(name=FONT, size=14, bold=True, color=NAVY)
for i, (k, v) in enumerate(notes, start=3):
    ws_n[f"A{i}"] = k; ws_n[f"B{i}"] = v
    ws_n[f"A{i}"].font = Font(name=FONT, bold=True, size=10)
    ws_n[f"B{i}"].font = Font(name=FONT, size=10)
    ws_n[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws_n[f"A{i}"].alignment = Alignment(vertical="top")

wb.save(OUT)
print("saved", OUT, "rows", N)
